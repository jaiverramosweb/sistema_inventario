"""
Customize pandoc's reference.docx with SI Sistemas branding.
- Header: project name + organization
- Footer: NIT + page number
- Heading colors: corporate blue (#1F3864)
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO_ROOT = Path(__file__).resolve().parents[2]
SRC = REPO_ROOT / "entrega_cliente" / "recursos" / "reference-default.docx"
OUT = REPO_ROOT / "entrega_cliente" / "recursos" / "reference.docx"

CORPORATE_BLUE = RGBColor(0x1F, 0x38, 0x64)
CORPORATE_DARK_GREY = RGBColor(0x33, 0x33, 0x33)

ORG_NAME = "SI SISTEMAS INFORMATICOS Y TECNOLOGIA SAS"
ORG_NIT = "NIT 900.583.147-1"
PROJECT_NAME = "Sistema de Inventario Pro"

doc = Document(str(SRC))

# Heading colors
HEADING_COLORS = {
    "Heading 1": (CORPORATE_BLUE, 18, True),
    "Heading 2": (CORPORATE_BLUE, 14, True),
    "Heading 3": (CORPORATE_BLUE, 12, True),
    "Heading 4": (CORPORATE_DARK_GREY, 11, True),
    "Heading 5": (CORPORATE_DARK_GREY, 11, False),
    "Heading 6": (CORPORATE_DARK_GREY, 10, False),
    "Title": (CORPORATE_BLUE, 24, True),
}

for style_name, (color, size, bold) in HEADING_COLORS.items():
    try:
        style = doc.styles[style_name]
        font = style.font
        font.color.rgb = color
        font.size = Pt(size)
        font.bold = bold
        font.name = "Calibri"
    except KeyError:
        pass

# Body
try:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
except KeyError:
    pass

# Header / footer on default section
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)
section.header_distance = Cm(1.0)
section.footer_distance = Cm(1.0)

# Header
header = section.header
header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
# Clear existing runs
for run in list(header_para.runs):
    run.text = ""
run = header_para.add_run(f"{PROJECT_NAME}  |  {ORG_NAME}")
run.font.name = "Calibri"
run.font.size = Pt(9)
run.font.color.rgb = CORPORATE_BLUE
run.font.bold = True

# Border bottom on header
pPr = header_para._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "6")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), "1F3864")
pBdr.append(bottom)
pPr.append(pBdr)

# Footer with PAGE field
footer = section.footer
footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in list(footer_para.runs):
    run.text = ""

run_left = footer_para.add_run(f"{ORG_NIT}    |    Confidencial    |    ")
run_left.font.name = "Calibri"
run_left.font.size = Pt(9)
run_left.font.color.rgb = CORPORATE_DARK_GREY

# Insert PAGE field
def add_page_field(paragraph):
    run = paragraph.add_run()
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = CORPORATE_DARK_GREY
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)

add_page_field(footer_para)

run_of = footer_para.add_run(" de ")
run_of.font.name = "Calibri"
run_of.font.size = Pt(9)
run_of.font.color.rgb = CORPORATE_DARK_GREY

def add_numpages_field(paragraph):
    run = paragraph.add_run()
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = CORPORATE_DARK_GREY
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "NUMPAGES"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)

add_numpages_field(footer_para)

# Border top on footer
pPr = footer_para._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
top = OxmlElement("w:top")
top.set(qn("w:val"), "single")
top.set(qn("w:sz"), "6")
top.set(qn("w:space"), "1")
top.set(qn("w:color"), "1F3864")
pBdr.append(top)
pPr.append(pBdr)

doc.save(str(OUT))
print(f"Reference saved: {OUT}")
